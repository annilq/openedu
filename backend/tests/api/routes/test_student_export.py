"""账号表导出端点（teacher-scale-up ticket 07 / ADR-0068 §2.2）。

覆盖验收点：
- 导出 xlsx 含班级 / 姓名 / 学号 / 初始密码四列
- 未分班学生班级字段标记「未分班」
- 导出的初始密码与导入时设定的配置值一致（= settings.STUDENT_DEFAULT_PASSWORD）
- 导出与导入共用同一 xlsx 库，且含已导入的学生

⚠️ 测试库 `db` 是 session 级、仅在整轮收尾清空，用户名须全文件唯一。
"""

import io
import uuid

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.config import settings
from app.db.models import User
from tests.utils.user import auth_headers, login, register_teacher

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _teacher(client: TestClient, username: str) -> tuple[str, str]:
    register_teacher(client, username=username, password="pw123456")
    token = login(client, username, "pw123456").json()["access_token"]
    me = client.get("/api/v1/auth/me", headers=auth_headers(token)).json()
    return token, me["id"]


def _make_student(client: TestClient, token: str, username: str, display_name: str) -> dict:
    r = client.post(
        "/api/v1/students",
        headers=auth_headers(token),
        json={
            "username": username,
            "password": "kid123456",
            "display_name": display_name,
            "grade": 2,
            "role": "student",
        },
    )
    assert r.status_code == 201, r.text
    return r.json()


def _build_xlsx(headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_export(content: bytes) -> list[tuple]:
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    return list(ws.iter_rows(values_only=True))


def test_export_columns_unassigned_and_password(client: TestClient, db):
    token, _ = _teacher(client, "exp_owner")
    cls = client.post(
        "/api/v1/classes",
        headers=auth_headers(token),
        json={"name": "一班", "grade": 3},
    )
    assert cls.status_code == 201, cls.text
    cid = cls.json()["id"]

    s1 = _make_student(client, token, "3101001", "张三")  # 会分到一班
    _make_student(client, token, "3101002", "李四")  # 保持未分班
    stu1 = db.get(User, uuid.UUID(s1["id"]))
    stu1.class_id = uuid.UUID(cid)
    db.add(stu1)
    db.commit()

    r = client.get("/api/v1/students/export", headers=auth_headers(token))
    assert r.status_code == 200, r.text
    assert "attachment" in r.headers.get("content-disposition", "")

    rows = _parse_export(r.content)
    header = list(rows[0])
    assert header == ["班级", "姓名", "学号", "初始密码"]
    data = rows[1:]
    by_no = {row[2]: row for row in data}

    assert by_no["3101001"][0] == "一班"  # 已分班学生显示班级名
    assert by_no["3101002"][0] == "未分班"  # 未分班学生标记明确
    # 初始密码与配置值一致
    assert by_no["3101001"][3] == settings.STUDENT_DEFAULT_PASSWORD
    assert by_no["3101002"][3] == settings.STUDENT_DEFAULT_PASSWORD


def test_export_includes_imported_students(client: TestClient):
    token, _ = _teacher(client, "exp_import")
    rows = [["王五", "3201001"], ["赵六", "3201002"]]
    data = _build_xlsx(["姓名", "学号"], rows)
    imp = client.post(
        "/api/v1/students/import",
        headers=auth_headers(token),
        files={"file": ("roster.xlsx", data, XLSX_CT)},
    )
    assert imp.status_code == 200
    assert imp.json()["created"] == 2

    r = client.get("/api/v1/students/export", headers=auth_headers(token))
    assert r.status_code == 200, r.text
    export_rows = _parse_export(r.content)
    by_no = {row[2]: row for row in export_rows[1:]}

    # 已导入学生在导出表中出现，且密码与导入配置一致（共用同一 xlsx 库）
    assert "3201001" in by_no
    assert by_no["3201001"][1] == "王五"
    assert by_no["3201001"][3] == settings.STUDENT_DEFAULT_PASSWORD
    assert by_no["3201002"][3] == settings.STUDENT_DEFAULT_PASSWORD
