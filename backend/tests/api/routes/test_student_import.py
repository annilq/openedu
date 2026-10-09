"""Excel 批量导入学生端点（teacher-scale-up ticket 05 / ADR-0068 §2.2）。

覆盖验收点：
- 正常导入创建账号，username=学号，初始密码=配置值，且能真正登录
- 缺列整行拒绝并回传行号（不猜测、不半成功）
- 学号冲突落 skipped，不覆盖既有账号
- 同文件重复导入全部 skipped（幂等）
- 大班（≥60 行）导入成功
- 文件超限（大小 / 行数 / 缺必要列）被拒

⚠️ 测试库 `db` 是 session 级、仅在整轮收尾清空，故各测试用户名必须全文件唯一，
   不能跨测试复用（同文件内前一个测试创建的行会持续存在）。
"""

import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.config import settings
from tests.utils.user import auth_headers, login, register_teacher

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _teacher(client: TestClient, username: str) -> tuple[str, object]:
    register_teacher(client, username=username, password="pw123456")
    token = login(client, username, "pw123456").json()["access_token"]
    me = client.get("/api/v1/auth/me", headers=auth_headers(token)).json()
    return token, me["id"]


def _make_student(client: TestClient, token: str, username: str, display_name: str) -> dict:
    r = client.post(
        "/api/v1/students",
        headers=auth_headers(token),
        json={
            "username": username,
            "password": "kid123456",
            "display_name": display_name,
            "grade": 2,
            "role": "student",
        },
    )
    assert r.status_code == 201, r.text
    return r.json()


def _build_xlsx(headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import(client: TestClient, token: str, data: bytes, filename: str = "roster.xlsx"):
    return client.post(
        "/api/v1/students/import",
        headers=auth_headers(token),
        files={"file": (filename, data, XLSX_CT)},
    )


def test_normal_import_creates_accounts_and_can_login(client: TestClient):
    token, _ = _teacher(client, "imp_owner")
    rows = [["张三", "2101001"], ["李四", "2101002"], ["王五", "2101003"]]
    data = _build_xlsx(["姓名", "学号"], rows)

    r = _import(client, token, data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["created"] == 3
    assert body["skipped"] == 0
    assert body["errors"] == []

    # 导入后该 username + 初始密码能真正登录
    for _, no in rows:
        login_r = client.post(
            "/api/v1/auth/login",
            json={"username": no, "password": settings.STUDENT_DEFAULT_PASSWORD},
        )
        assert login_r.status_code == 200, login_r.text
        assert login_r.json()["access_token"]


def test_missing_column_returns_row_number(client: TestClient):
    token, _ = _teacher(client, "imp_missing")
    rows = [
        ["张三", "2120001"],  # 正常
        ["", "2120002"],  # 缺姓名
        ["李四", ""],  # 缺学号
        ["王五", "2120003"],  # 正常
    ]
    data = _build_xlsx(["姓名", "学号"], rows)

    r = _import(client, token, data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["created"] == 2
    assert body["skipped"] == 2
    reasons = {e["row"]: e["reason"] for e in body["errors"]}
    # Excel 行号：表头是第 1 行，故第 2 条数据（缺姓名）是第 3 行、第 3 条（缺学号）是第 4 行。
    assert reasons[3] == "缺少姓名"
    assert reasons[4] == "缺少学号"


def test_duplicate_student_no_skipped_not_overwritten(client: TestClient):
    token, _ = _teacher(client, "imp_dup")
    # 预先存在一个学号相同的老账号（display_name=旧名）
    _make_student(client, token, "2150001", "旧名")
    rows = [
        ["新名", "2150001"],  # 学号已存在 → skipped，不覆盖
        ["赵六", "2150099"],  # 正常
    ]
    data = _build_xlsx(["姓名", "学号"], rows)

    r = _import(client, token, data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["created"] == 1
    assert body["skipped"] == 1
    assert body["errors"][0]["reason"] == "学号已存在"

    # 既有账号 display_name 未被覆盖
    lst = client.get("/api/v1/students", headers=auth_headers(token)).json()
    matched = [u for u in lst["data"] if u["username"] == "2150001"]
    assert matched and matched[0]["display_name"] == "旧名"


def test_idempotent_reimport_all_skipped(client: TestClient):
    token, _ = _teacher(client, "imp_idem")
    rows = [["张三", "2166001"], ["李四", "2166002"]]
    data = _build_xlsx(["姓名", "学号"], rows)

    first = _import(client, token, data)
    assert first.status_code == 200
    assert first.json()["created"] == 2

    second = _import(client, token, data)
    assert second.status_code == 200, second.text
    body = second.json()
    assert body["created"] == 0
    assert body["skipped"] == 2
    assert all(e["reason"] == "学号已存在" for e in body["errors"])


def test_large_class_ge60(client: TestClient):
    token, _ = _teacher(client, "imp_big")
    rows = [[f"学生{i:02d}", f"2099{i:04d}"] for i in range(65)]  # 65 行
    assert len(rows) >= 60
    data = _build_xlsx(["姓名", "学号"], rows)

    r = _import(client, token, data)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["created"] == 65
    assert body["skipped"] == 0
    assert body["errors"] == []


def test_file_too_large_rejected(client: TestClient):
    token, _ = _teacher(client, "imp_size")
    data = b"a" * (settings.STUDENT_IMPORT_MAX_BYTES + 1)

    r = _import(client, token, data, filename="huge.xlsx")
    assert r.status_code == 413, r.text


def test_too_many_rows_rejected(client: TestClient):
    token, _ = _teacher(client, "imp_rows")
    rows = [[f"学生{i}", f"3000{i:04d}"] for i in range(settings.STUDENT_IMPORT_MAX_ROWS + 1)]
    data = _build_xlsx(["姓名", "学号"], rows)

    r = _import(client, token, data)
    assert r.status_code == 400, r.text
    assert "行数上限" in r.json()["message"]


def test_missing_required_header_rejected(client: TestClient):
    token, _ = _teacher(client, "imp_header")
    # 缺「学号」列
    data = _build_xlsx(["姓名", "班级"], [["张三", "一班"]])

    r = _import(client, token, data)
    assert r.status_code == 400, r.text
    assert "缺少必要列" in r.json()["message"]


def test_import_template_download(client: TestClient):
    """方案A：下载导入模板返回 xlsx，表头与导入解析口径一致（姓名 / 学号）。

    模板不含数据行，避免教师直接上传示例数据产生脏账号。
    """
    token, _ = _teacher(client, "imp_tpl")
    r = client.get(
        "/api/v1/students/import-template",
        headers=auth_headers(token),
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == XLSX_CT
    assert "student_import_template.xlsx" in r.headers.get("content-disposition", "")

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows == [("姓名", "学号")], rows
