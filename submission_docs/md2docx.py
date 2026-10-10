#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 submission_docs/ 下的 Markdown 手册转成 .docx（复用 gen.py 的中文字体排版助手）。

单一事实源：Markdown 文件（使用手册.md / 安装手册.md / 开发记录.md）。
改完 Markdown 后重跑本脚本即可刷新对应 .docx，不用再改两处。

用法：
  /Users/annilq/.workbuddy/binaries/python/envs/default/bin/python md2docx.py
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH

from PIL import Image

# 复用 gen.py 的排版助手与字体常量（gen.py 有 __main__ guard，import 安全）
from gen import (
    add_title, add_h1, add_h2, add_h3, add_para, add_code,
    set_cell, set_col_widths, set_run_font, set_first_indent,
    BODY, H1, H2, H3, CODE, LINE,
)

HERE = Path(__file__).resolve().parent

# python-docx 默认模板是 Letter 纸（左右各 1.25"），可用宽仅 15.24cm
# 表格列宽合计超过就会溢出页面，这里钉死总宽
USABLE_CM = 15.24


def add_rich(doc, text, font=BODY, size=16, indent=False, left_indent=None,
             space_after=6, line=LINE):
    """支持 **加粗** 与 `行内代码` 的段落。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line)
    pf.space_after = Pt(space_after)
    if left_indent is not None:
        pf.left_indent = Pt(left_indent)
    if indent:
        set_first_indent(p)
    # 先切加粗，再在非加粗片段里切行内代码
    # （顺序必须如此：否则 **`code`** 这种嵌套写法会被拆坏）
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            inner = part[2:-2].replace("`", "")  # 加粗片段内的反引号去掉
            run = p.add_run(inner)
            set_run_font(run, font, size, bold=True)
            continue
        for seg in re.split(r"(`[^`]+`)", part):
            if not seg:
                continue
            if seg.startswith("`") and seg.endswith("`"):
                run = p.add_run(seg[1:-1])
                set_run_font(run, CODE, size)
            else:
                run = p.add_run(seg)
                set_run_font(run, font, size)
    return p


def col_widths(ncols: int):
    """按可用宽度 15.24cm 分配列宽；两列表格首列窄一些更好读。"""
    if ncols == 2:
        return [Cm(4.5), Cm(USABLE_CM - 4.5)]
    w = USABLE_CM / ncols
    return [Cm(w) for _ in range(ncols)]


def flush_table(doc, rows):
    """rows: list[list[str]]，首行当表头。"""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncols):
            # 单元格走 set_cell（不解析富文本），先把 ** / ` 标记剥掉，避免原样输出
            text = (row[j] if j < len(row) else "").replace("**", "").replace("`", "")
            set_cell(table.cell(i, j), text, font=BODY, size=10.5,
                     bold=(i == 0), first=True)
    set_col_widths(table, col_widths(ncols))
    add_para(doc, "", size=6, line=8, space_after=0)


def add_image(doc, caption: str, rel_path: str) -> None:
    """嵌入图片 + 居中图注。宽屏图用满可用宽，竖长截图收窄以免超页高。"""
    img = HERE / rel_path
    if not img.exists():
        add_rich(doc, f"（缺图：{rel_path}）")
        return
    with Image.open(img) as im:
        aspect = im.width / im.height
    width_cm = 15.0 if aspect >= 1.2 else 11.0
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(img), width=Cm(width_cm))
    if caption:
        add_para(doc, caption, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_after=12)


def convert(md_path: Path, docx_path: Path) -> Path:
    doc = Document()
    lines = md_path.read_text(encoding="utf-8").splitlines()

    in_code = False
    table_rows: list[list[str]] = []

    def end_table():
        nonlocal table_rows
        if table_rows:
            flush_table(doc, table_rows)
            table_rows = []

    for raw in lines:
        line = raw.rstrip()

        # 围栏代码块
        if line.startswith("```"):
            if in_code:
                in_code = False
            else:
                end_table()
                in_code = True
            continue
        if in_code:
            add_code(doc, line)
            continue

        stripped = line.strip()

        if not stripped:
            end_table()
            continue

        # 表格行
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue  # 分隔行
            table_rows.append(cells)
            continue
        end_table()

        # 水平分割线
        if re.fullmatch(r"-{3,}", stripped):
            continue

        # 标题
        if stripped.startswith("# "):
            add_title(doc, stripped[2:].strip())
            continue
        if stripped.startswith("## "):
            add_h1(doc, stripped[3:].strip())
            continue
        if stripped.startswith("### "):
            add_h2(doc, stripped[4:].strip())
            continue
        if stripped.startswith("#### "):
            add_h3(doc, stripped[5:].strip())
            continue

        # 引用
        if stripped.startswith(">"):
            add_rich(doc, stripped.lstrip("> ").strip(), size=16)
            continue

        # 编号列表
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            add_rich(doc, f"{m.group(1)}. {m.group(2)}", left_indent=18)
            continue

        # 无序列表
        if stripped.startswith(("- ", "* ")):
            add_rich(doc, "• " + stripped[2:], left_indent=18)
            continue

        # 图片：![图注](相对路径)
        m = re.match(r"^!\[(.*?)\]\((.+?)\)$", stripped)
        if m:
            add_image(doc, m.group(1).strip(), m.group(2).strip())
            continue

        # 正文
        add_rich(doc, stripped, indent=True)

    end_table()
    doc.save(str(docx_path))
    return docx_path


def main() -> None:
    pairs = [
        ("使用手册.md", "使用手册.docx"),
        ("安装手册.md", "安装手册.docx"),
        ("开发与应用报告.md", "开发与应用报告.docx"),
        ("开发记录.md", "开发记录.docx"),
        ("提示词开发流程与提示词集.md", "提示词开发流程与提示词集.docx"),
    ]
    for src, dst in pairs:
        md = HERE / src
        if not md.exists():
            print(f"[md2docx] 跳过（缺 {src}）")
            continue
        out = convert(md, HERE / dst)
        print(f"[md2docx] {src} -> {out.name}")


if __name__ == "__main__":
    sys.exit(main())
